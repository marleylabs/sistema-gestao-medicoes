"""Diagnóstico SOMENTE LEITURA das identidades de uma planilha de medição (aba Documentos).

Uso:
  python diagnose_identities.py <planilha.xlsx|xlsm> [--sheet Documentos] [--database-url URL] [--ciclo YYMM]

Para cada linha classifica, com a MESMA regra da importação:
  RESOLVED          nome resolve por código/alias/legado
  AUTO_MATCH        correspondência determinística única (seria gravada como alias na importação)
  PENDING           nome plausível sem vínculo (importa, fica no ciclo, BM bloqueado)
  STRUCTURAL_ERROR  linha com dados sem PROJETISTA, ou PROJETISTA que não é nome (bloqueia)
  DISCARDED         nome/linha já descartado no ciclo (só com --database-url e --ciclo)
  EMPTY_IGNORED     linha sem chave de medição nem desconto (nunca conta)
Também mostra o valor bruto de PROJETISTA, se a célula está num intervalo mesclado e o valor
efetivo usado. Sem --database-url (ou ETL_DATABASE_URL) não consulta cadastros: todo nome
plausível aparece como PENDING. Nunca grava no banco (só `engine.connect()` de leitura) e nunca
altera a planilha.
"""
from __future__ import annotations

import argparse
import os
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from ingest_medicoes import (
    DEFAULT_SHEET,
    IDENTITY_HEADER,
    MEASUREMENT_COLUMNS,
    PROFESSIONAL_COLUMNS,
    PROJECT_COLUMNS,
    SHEET_ALIASES,
    OperationalIdentityResolver,
    clean_text,
    extract,
    first_value,
    has_discount_data,
    is_valid_measurement_key,
    load_import_decisions,
    looks_like_supplier_name,
    normalize_person_name,
    read_measurements_sheet,
    resolve_sheet_name,
    row_decision_key,
    should_keep_vba,
)


def carregar_contexto(database_url: str | None, ciclo: str | None):
    if not database_url:
        return OperationalIdentityResolver([], []), {}
    from sqlalchemy import create_engine

    engine = create_engine(database_url)
    with engine.connect() as conn:  # leitura apenas
        resolver = OperationalIdentityResolver.load(conn)
        decisoes = load_import_decisions(conn, ciclo) if ciclo else {}
    return resolver, decisoes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("planilha", type=Path)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--database-url", default=os.environ.get("ETL_DATABASE_URL"))
    parser.add_argument("--ciclo", default=None, help="Ciclo (YYMM) para considerar descartes já registrados.")
    args = parser.parse_args()

    sheet_name = resolve_sheet_name(args.planilha, args.sheet, SHEET_ALIASES["Documentos"])
    df = read_measurements_sheet(args.planilha, sheet_name)
    excel_rows = df.attrs.get("excel_row_numbers", [])
    resolver, decisoes = carregar_contexto(args.database_url, args.ciclo)
    nomes_descartados = set((decisoes.get("nomes_descartados") or {}).keys())
    linhas_descartadas = set(decisoes.get("linhas_descartadas") or ())

    linhas = []
    for position, (_index, row) in enumerate(df.iterrows()):
        project = extract(row, PROJECT_COLUMNS)
        numero = clean_text(first_value(row, MEASUREMENT_COLUMNS["numero_medicao"]))
        relevante = is_valid_measurement_key(numero, clean_text(project["codigo_projeto"])) or has_discount_data(row)
        linha = excel_rows[position] if position < len(excel_rows) else None
        linhas.append((linha, row, relevante, clean_text(extract(row, PROFESSIONAL_COLUMNS)["nome"])))
    # Mesma ordem da importação: correspondências automáticas avaliadas sobre todos os nomes da planilha.
    resolver.plan_auto_matches([nome for _l, _r, relevante, nome in linhas if relevante and nome])

    workbook = load_workbook(args.planilha, read_only=False, data_only=True, keep_vba=should_keep_vba(args.planilha))
    try:
        sheet = workbook[sheet_name]
        header_cell = next(
            (cell for row in sheet.iter_rows(max_row=100) for cell in row if clean_text(cell.value) == IDENTITY_HEADER),
            None,
        )
        column = header_cell.column if header_cell else None
        merged = [r for r in sheet.merged_cells.ranges if column and r.min_col <= column <= r.max_col]
        print(f"Aba: {sheet_name} · coluna {IDENTITY_HEADER}: {header_cell.coordinate if header_cell else 'não encontrada'}")
        print(f"Intervalos mesclados na coluna: {len(merged)} · banco: {'sim (leitura)' if args.database_url else 'não'}")
        print("linha\tcategoria\tbruto\tmesclado\tefetivo\tdetalhe")

        contagem: Counter[str] = Counter()
        for linha, row, relevante, nome in linhas:
            bruto = clean_text(sheet.cell(linha, column).value) if linha and column else None
            faixa = next((r for r in merged if linha and r.min_row <= linha <= r.max_row), None)
            detalhe = ""
            if not relevante:
                categoria = "EMPTY_IGNORED"
            elif not nome:
                if args.ciclo and row_decision_key(row, args.ciclo) in linhas_descartadas:
                    categoria = "DISCARDED"
                else:
                    categoria = "STRUCTURAL_ERROR"
                    detalhe = f"sem PROJETISTA · documento {clean_text(first_value(row, MEASUREMENT_COLUMNS['numero_documento'])) or '—'}"
            elif normalize_person_name(nome) in nomes_descartados:
                categoria = "DISCARDED"
            else:
                resultado = resolver.resolve(nome)
                if resultado["status"] == "RESOLVIDO":
                    categoria = "AUTO_MATCH" if resultado["via"] == "AUTO_MATCH" else "RESOLVED"
                    detalhe = f'→ {resultado["codigo"]} ({resultado["via"]})'
                elif resultado["status"] == "NAO_RESOLVIDO" and looks_like_supplier_name(nome):
                    categoria = "PENDING"
                else:
                    categoria = "STRUCTURAL_ERROR"
                    detalhe = "PROJETISTA não é nome de fornecedor" if resultado["status"] == "NAO_RESOLVIDO" else resultado["status"]
            contagem[categoria] += 1
            if categoria != "EMPTY_IGNORED":
                print(f"{linha}\t{categoria}\t{bruto or ''}\t{faixa.coord if faixa else ''}\t{nome or '(vazio)'}\t{detalhe}")

        print("\nResumo:")
        for categoria in ("RESOLVED", "AUTO_MATCH", "PENDING", "STRUCTURAL_ERROR", "DISCARDED", "EMPTY_IGNORED"):
            print(f"  {categoria}: {contagem[categoria]}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
