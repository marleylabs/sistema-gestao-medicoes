"""Identidades pendentes, correspondência automática e descarte na importação (fixture sintética,
sem nenhum nome/documento real).

- linha realmente vazia (sem dado, só formatação, só espaços) nunca é ocorrência/erro/pendência;
- PROJETISTA mesclado: linhas do intervalo herdam a âncora antes da avaliação de "vazio";
- linha com dados sem PROJETISTA = erro estrutural (bloqueia) até ser descartada conscientemente;
- AUTO_MATCH só quando determinístico e único ("MARIA COSTA" → "MARIA FERNANDA LIMA COSTA"),
  nunca com 2 candidatos ("JOAO SILVA"), nunca por distância, nunca com nomes que colidem;
- nome plausível sem cadastro ("CARLOS RAMOS") vira PENDENTE e não bloqueia;
- descarte é do ciclo; carga real no Postgres E2E (transação) cobre atomicidade, pendente sem
  Profissional, mapa gerado, alias do AUTO_MATCH e reimportação idempotente."""
from __future__ import annotations

import os
import tempfile
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

import ingest_medicoes as im
from ingest_medicoes import (
    OperationalIdentityResolver,
    UnresolvedIdentityError,
    assert_operational_identities_resolved,
    looks_like_supplier_name,
    read_measurements_sheet,
    row_decision_key,
)
from test_identity_extraction import HEADERS, row_values
from test_negative_measurement_validation import normal_row

MARIA = {"id": "p-maria", "codigo": "MARIA FERNANDA LIMA COSTA", "nome": "MARIA FERNANDA LIMA COSTA"}
JOAO_A = {"id": "p-joao-a", "codigo": "JOAO PEDRO SILVA", "nome": "JOAO PEDRO SILVA"}
JOAO_B = {"id": "p-joao-b", "codigo": "JOAO CARLOS SILVA", "nome": "JOAO CARLOS SILVA"}
ITALO = {"id": "p-italo", "codigo": "ITALO RUAN BRAGA VIANA", "nome": "ITALO RUAN BRAGA VIANA"}
ALFA = {"id": "p-alfa", "codigo": "FORNECEDOR ALFA TESTE", "nome": "FORNECEDOR ALFA TESTE"}
PROFISSIONAIS = [MARIA, JOAO_A, JOAO_B, ITALO, ALFA]
CANDIDATOS = [{"id": p["id"], "codigo": p["codigo"], "responsavel": p["codigo"]} for p in PROFISSIONAIS]


def resolver() -> OperationalIdentityResolver:
    return OperationalIdentityResolver(PROFISSIONAIS, [], CANDIDATOS)


def preflight(df: pd.DataFrame, r: OperationalIdentityResolver | None = None, decisoes=None, bm_aux: pd.DataFrame | None = None):
    return assert_operational_identities_resolved(
        df, pd.DataFrame(), bm_aux if bm_aux is not None else pd.DataFrame(), pd.DataFrame(), set(), {}, {},
        "2804", r or resolver(), {}, "Documentos", None, "Documentos Auxiliares", None, decisoes,
    )


# ─── linhas vazias e merge ───────────────────────────────────────────────────────────────────
def build_empty_rows_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(HEADERS)                                              # 1
    ws.append(row_values(PROJETISTA="FORNECEDOR ALFA TESTE"))        # 2
    col = HEADERS.index("PROJETISTA") + 1
    for r in range(3, 6):                                           # 3–5: PROJETISTA mesclado, com documento
        ws.append(row_values(PROJETISTA=None, **{"Número do Documento": f"DOC-MESCLA-{r}"}))
    ws.cell(3, col).value = "FORNECEDOR ALFA TESTE"
    ws.merge_cells(start_row=3, start_column=col, end_row=5, end_column=col)
    ws.append([None] * len(HEADERS))                                # 6: vazia
    ws.append(["   "] * len(HEADERS))                               # 7: só espaços
    for c in range(1, len(HEADERS) + 1):                            # 8: só formatação
        ws.cell(8, c).fill = PatternFill("solid", fgColor="FFFF00")
    ws.cell(100, 1).fill = PatternFill("solid", fgColor="FFFF00")   # 100: só formatação, bem abaixo
    wb.save(path)


def test_empty_rows_are_ignored_and_merged_rows_inherit() -> None:
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "vazias.xlsx"
        build_empty_rows_workbook(p)
        df = read_measurements_sheet(p, "Documentos")
    assert df["PROJETISTA"].tolist() == ["FORNECEDOR ALFA TESTE"] * 4  # 2 + 3 linhas do merge; vazias fora
    summary = preflight(df)  # nada pendente, nada bloqueado, nenhuma ocorrência vazia
    assert summary["identidades_resolvidas"] == 1 and summary["identidades_pendentes"] == []
    assert summary["linhas_descartadas"] == 0


def test_whitespace_projetista_on_data_row_is_structural() -> None:
    df = pd.DataFrame([normal_row(PROJETISTA="   ", Evidência="GRD-T-SINT-0001")])
    df.attrs["excel_row_numbers"] = [40]
    try:
        preflight(df)
    except UnresolvedIdentityError as error:
        assert error.details[0]["status"] == "SEM_PROJETISTA" and error.details[0]["linhas"] == [40]
        assert not any(d["valor"].startswith("GRD") for d in error.details)  # GRD da Evidência nunca vira nome
    else:
        raise AssertionError("Linha com dados e PROJETISTA vazio deveria bloquear.")


# ─── AUTO_MATCH ──────────────────────────────────────────────────────────────────────────────
def test_auto_match_only_when_deterministic_and_unique() -> None:
    r = resolver()
    nomes = ["MARIA COSTA", "JOAO SILVA", "CARLOS RAMOS", "MARIA FERNANDA LIMA COSTA", "MARYA COSTA"]
    plano = r.plan_auto_matches(nomes)
    assert set(plano) == {"MARIA COSTA"}                        # tokens exatos, mesmo 1º/último, candidato único
    assert plano["MARIA COSTA"]["codigo"] == "MARIA FERNANDA LIMA COSTA"
    assert r.resolve("maria costa")["via"] == "AUTO_MATCH"
    assert r.resolve("MARIA FERNANDA LIMA COSTA")["via"] == "CODIGO"  # canônico exato continua CODIGO
    assert r.resolve("JOAO SILVA")["status"] == "NAO_RESOLVIDO"       # dois candidatos
    assert r.resolve("MARYA COSTA")["status"] == "NAO_RESOLVIDO"      # parecido (fuzzy) nunca
    assert r.resolve("MARIA")["status"] == "NAO_RESOLVIDO"            # um token só
    # Um "MARIA" solto na mesma planilha também cabe no cadastro → colisão → nada automático.
    assert resolver().plan_auto_matches(["MARIA COSTA", "MARIA"]) == {}


def test_alias_exact_wins_and_colliding_names_stay_manual() -> None:
    r = OperationalIdentityResolver(PROFISSIONAIS, [{"profissional_id": "p-alfa", "alias": "ALFA APELIDO", "alias_normalizado": "ALFA APELIDO"}], CANDIDATOS)
    assert r.plan_auto_matches(["ALFA APELIDO"]) == {}               # já resolve por alias
    assert r.resolve("ALFA APELIDO")["via"] == "ALIAS"
    # "ITALO VIANA" cabe sozinho em ITALO RUAN BRAGA VIANA, mas "ITALO RUAN" também → manual.
    assert r.plan_auto_matches(["ITALO VIANA", "ITALO RUAN"]) == {}
    assert OperationalIdentityResolver(PROFISSIONAIS, [], CANDIDATOS).plan_auto_matches(["ITALO VIANA"]).keys() == {"ITALO VIANA"}
    # Nunca com cadastro inativo/sem Profissional: só `candidatos` entram.
    assert OperationalIdentityResolver(PROFISSIONAIS, [], []).plan_auto_matches(["MARIA COSTA"]) == {}


def test_preflight_classifies_resolved_auto_pending() -> None:
    rows = [normal_row(PROJETISTA=n) for n in ("FORNECEDOR ALFA TESTE", "MARIA COSTA", "MARIA COSTA", "JOAO SILVA", "CARLOS RAMOS")]
    df = pd.DataFrame(rows)
    df.attrs["excel_row_numbers"] = [5, 6, 7, 8, 9]
    summary = preflight(df)
    assert summary["correspondencias_automaticas"] == ["MARIA COSTA -> MARIA FERNANDA LIMA COSTA"]
    assert {p["valor"] for p in summary["identidades_pendentes"]} == {"JOAO SILVA", "CARLOS RAMOS"}
    assert summary["plano"]["auto"]["MARIA COSTA"]["ocorrencias"] == 2
    assert summary["identidades_por_via"] == {"CODIGO": 1, "AUTO_MATCH": 1}


def test_parser_residue_in_projetista_is_structural_never_pending() -> None:
    for lixo in ("GRD-T-SINT-2026-0001-0001", "HORAS DE ESTUDO PARA ADEQUACAO DO PROJETO DE TUBULACAO DA AREA SINTETICA"):
        df = pd.DataFrame([normal_row(PROJETISTA=lixo)])
        try:
            preflight(df)
        except UnresolvedIdentityError as error:
            assert error.details[0]["status"] == "PROJETISTA_INVALIDO"
        else:
            raise AssertionError(f"{lixo!r} não pode virar fornecedor pendente.")
    assert looks_like_supplier_name("TUBULAÇÃO") is True  # palavra única é aceita como nome; a coluna de disciplina nunca é lida como identidade
    # Dígito isolado é nome legítimo de empresa (nunca bloqueia a importação por isso).
    for empresa in ("A1 ENGENHARIA SINTETICA", "3D PROJETOS SINT", "ENGENHARIA 360 SINT"):
        assert looks_like_supplier_name(empresa) is True, empresa


# ─── decisões de descarte ────────────────────────────────────────────────────────────────────
def test_discarded_name_and_row_are_skipped_and_do_not_block() -> None:
    sem = normal_row(PROJETISTA=None, **{"Número do Documento": "DOC-SEM-1"})
    rows = [normal_row(PROJETISTA="CARLOS RAMOS"), sem, normal_row(PROJETISTA="FORNECEDOR ALFA TESTE")]
    df = pd.DataFrame(rows)
    chave = row_decision_key(df.iloc[1], "2804")
    decisoes = {"nomes_descartados": {"CARLOS RAMOS": "CARLOS RAMOS"}, "linhas_descartadas": {chave}, "existentes": {}}
    summary = preflight(df, decisoes=decisoes)
    assert summary["identidades_pendentes"] == []
    assert summary["identidades_descartadas"] == [{"valor": "CARLOS RAMOS", "ocorrencias": 1}]
    assert summary["linhas_descartadas"] == 1
    # Sem a decisão (outro ciclo): a mesma linha bloqueia e o nome volta a ser pendente.
    try:
        preflight(df)
    except UnresolvedIdentityError as error:
        assert error.details[0]["chaves"][0]["chave"] == chave
    else:
        raise AssertionError("Sem descarte, a linha sem PROJETISTA bloqueia.")
    assert row_decision_key(df.iloc[1], "2805") != chave  # descarte é do ciclo


# ─── carga real no Postgres E2E (transação com rollback) ─────────────────────────────────────
def _workbook_carga(path: Path, sufixo: str, incluir_sem_projetista: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(HEADERS)
    ws.append(row_values(PROJETISTA=f"FORNECEDOR RESOLVIDO {sufixo}", **{"Número do Documento": "DOC-A", "VALOR DE MEDIÇÃO": 100}))
    ws.append(row_values(PROJETISTA=f"CARLOS RAMOS {sufixo}", **{"Número do Documento": "DOC-B1", "VALOR DE MEDIÇÃO": 30}))
    ws.append(row_values(PROJETISTA=f"CARLOS RAMOS {sufixo}", **{"Número do Documento": "DOC-B2", "VALOR DE MEDIÇÃO": 20}))
    ws.append(row_values(PROJETISTA=f"MARIA {sufixo} COSTA", **{"Número do Documento": "DOC-C", "VALOR DE MEDIÇÃO": 10}))
    if incluir_sem_projetista:
        ws.append(row_values(PROJETISTA=None, **{"Número do Documento": "DOC-SEM", "VALOR DE MEDIÇÃO": 7}))
    wb.create_sheet("Documentos Auxiliares").append(["Responsavel", "Auxiliar", "Ciclo"])
    wb.save(path)


def test_real_load_pending_auto_match_and_atomicity() -> None:
    url_text = os.environ.get("ETL_DATABASE_URL") or os.environ.get("DATABASE_URL_TEST")
    if not url_text:
        print("SKIP: sem ETL_DATABASE_URL/DATABASE_URL_TEST")
        return
    from sqlalchemy import create_engine, text
    from sqlalchemy.engine import make_url

    url = make_url(url_text).difference_update_query(["schema"])
    assert url.host in ("localhost", "127.0.0.1") and "e2e" in (url.database or "")
    db = url.render_as_string(hide_password=False)
    engine = create_engine(db)
    s = uuid.uuid4().hex[:6].upper().replace("0", "X").replace("1", "Y").replace("2", "Z").replace("3", "W").replace("4", "K").replace("5", "J").replace("6", "Q").replace("7", "V").replace("8", "H").replace("9", "N")
    ciclo = "2807"
    resolvido = f"FORNECEDOR RESOLVIDO {s}"
    maria = f"MARIA {s} FERNANDA LIMA COSTA"
    ids: list[str] = []

    def contagens(conn):
        return {t: conn.execute(text(f"select count(*) from {t}")).scalar() for t in
                ("medicoes", "mapa_pagamento_itens", "profissional_aliases", "importacao_identidades", "profissionais", "admin_audit_logs")}

    with engine.begin() as conn:
        for codigo in (resolvido, maria):
            pid = conn.execute(text("insert into profissionais (nome, codigo) values (:c, :c) returning id"), {"c": codigo}).scalar()
            ids.append(str(pid))
            conn.execute(text("insert into cadastros_fornecedores (cnpj_normalizado, colaborador_codigo, responsavel, razao_social, raw_payload) values ('11222333000181', :c, :c, :c, '{}')"), {"c": codigo})
        conn.execute(text("delete from importacao_identidades where ciclo = :c"), {"c": ciclo})
    try:
        with tempfile.TemporaryDirectory() as d:
            bloqueada = Path(d) / "bloqueada.xlsx"
            _workbook_carga(bloqueada, s, incluir_sem_projetista=True)
            with engine.connect() as conn:
                antes = contagens(conn)
            try:
                im.ingest(bloqueada, "Documentos", "Base", "MAPA PAGTO", "Documentos Auxiliares", db, False, True, ciclo)
                raise AssertionError("Linha sem PROJETISTA deveria bloquear.")
            except UnresolvedIdentityError as error:
                chave_sem = error.details[0]["chaves"][0]["chave"]
            with engine.connect() as conn:
                assert contagens(conn) == antes, "bloqueio estrutural não pode gravar nada (nem alias do AUTO_MATCH)"

            # Descarte consciente da linha → a mesma planilha conclui.
            with engine.begin() as conn:
                conn.execute(text("insert into importacao_identidades (ciclo, tipo, chave, valor_bruto, origem, status) values (:c, 'LINHA_SEM_PROJETISTA', :k, '', 'Documentos', 'DESCARTADO')"), {"c": ciclo, "k": chave_sem})
            r1 = im.ingest(bloqueada, "Documentos", "Base", "MAPA PAGTO", "Documentos Auxiliares", db, False, True, ciclo,
                           importado_por=None)
            assert r1["correspondencias_automaticas"] == [f"MARIA {s} COSTA -> {maria}"]
            assert [p["valor"] for p in r1["identidades_pendentes"]] == [f"CARLOS RAMOS {s}"]
            assert r1["linhas_descartadas"] == 1 and "plano" not in r1

            with engine.connect() as conn:
                pend = conn.execute(text("select id, status, ocorrencias from importacao_identidades where ciclo = :c and tipo = 'IDENTIDADE' and chave = :k"),
                                    {"c": ciclo, "k": f"CARLOS RAMOS {s}"}).mappings().one()
                assert pend["status"] == "PENDENTE" and pend["ocorrencias"] == 2
                med = conn.execute(text("select count(*), sum(valor_medicao), count(id_profissional) from medicoes where ciclo = :c and identidade_importacao_id = :i"),
                                   {"c": ciclo, "i": pend["id"]}).one()
                assert med[0] == 2 and float(med[1]) == 50.0 and med[2] == 0  # no ciclo, com valor, sem Profissional
                mapa = conn.execute(text("select projetista_codigo, valor from mapa_pagamento_itens where ciclo = :c and identidade_importacao_id = :i"),
                                    {"c": ciclo, "i": pend["id"]}).one()
                assert mapa[0] == f"CARLOS RAMOS {s}" and float(mapa[1]) == 50.0
                assert conn.execute(text("select count(*) from profissionais where upper(nome) = :n or upper(codigo) = :n"), {"n": f"CARLOS RAMOS {s}"}).scalar() == 0
                alias = conn.execute(text("select origem from profissional_aliases where alias_normalizado = :n"), {"n": f"MARIA {s} COSTA"}).scalar()
                assert alias == "DOCUMENTOS"
                total1 = float(conn.execute(text("select sum(valor) from mapa_pagamento_itens where ciclo = :c"), {"c": ciclo}).scalar())
                assert total1 == 160.0  # 100 resolvido + 50 pendente + 10 auto; a linha descartada (7) não entra
                n1 = conn.execute(text("select count(*) from medicoes where ciclo = :c"), {"c": ciclo}).scalar()

            # Reimportação do MESMO arquivo: nada duplica; o AUTO_MATCH agora resolve por alias.
            r2 = im.ingest(bloqueada, "Documentos", "Base", "MAPA PAGTO", "Documentos Auxiliares", db, False, True, ciclo)
            assert r2["correspondencias_automaticas"] == [] and r2["identidades_por_via"].get("ALIAS") == 1
            with engine.connect() as conn:
                assert conn.execute(text("select count(*) from medicoes where ciclo = :c"), {"c": ciclo}).scalar() == n1
                assert float(conn.execute(text("select sum(valor) from mapa_pagamento_itens where ciclo = :c"), {"c": ciclo}).scalar()) == total1
                assert conn.execute(text("select count(*) from importacao_identidades where ciclo = :c and status = 'PENDENTE'"), {"c": ciclo}).scalar() == 1

            # Nome descartado neste ciclo: some do ciclo na reimportação; outro ciclo não é afetado.
            with engine.begin() as conn:
                conn.execute(text("update importacao_identidades set status = 'DESCARTADO' where id = :i"), {"i": pend["id"]})
            r3 = im.ingest(bloqueada, "Documentos", "Base", "MAPA PAGTO", "Documentos Auxiliares", db, False, True, ciclo)
            assert r3["identidades_descartadas"] == [{"valor": f"CARLOS RAMOS {s}", "ocorrencias": 2}] and r3["identidades_pendentes"] == []
    finally:
        with engine.begin() as conn:
            conn.execute(text("delete from medicoes where ciclo = :c"), {"c": ciclo})
            conn.execute(text("delete from mapa_pagamento_itens where ciclo = :c"), {"c": ciclo})
            conn.execute(text("delete from mapa_pagamento_contexto where ciclo = :c"), {"c": ciclo})
            conn.execute(text("delete from importacao_identidades where ciclo = :c"), {"c": ciclo})
            conn.execute(text("delete from etl_execucoes where ciclo = :c"), {"c": ciclo}) if conn.execute(text("select to_regclass('public.etl_execucoes') is not null")).scalar() else None
            conn.execute(text("delete from cadastros_fornecedores where colaborador_codigo in :c").bindparams(im.bindparam("c", expanding=True)), {"c": [resolvido, maria]})
            conn.execute(text("delete from profissionais where id in :ids").bindparams(im.bindparam("ids", expanding=True)), {"ids": ids})


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok {name}")
    print("OK: identidades pendentes, AUTO_MATCH, linhas vazias e descarte na importação.")
