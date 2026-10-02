"""Extração da identidade do fornecedor na aba Documentos (fixture sintética, sem dados reais).

Bug pós-deploy: com PROJETISTA vazio (célula mesclada ou linha sem projetista) o ETL caía para
"Evidência"/"Número do Documento" e números de GRD, descrições ("HORAS DE ESTUDO…", "TUBULAÇÃO")
e números de orçamento viravam "identidades não resolvidas". Agora:
- a identidade vem SÓ de PROJETISTA;
- célula mesclada de PROJETISTA herda o valor da âncora (só essa coluna);
- linha de medição sem PROJETISTA bloqueia como SEM_PROJETISTA (nunca vira alias/cadastro) até ser
  corrigida ou descartada conscientemente; nome desconhecido não bloqueia (vira pendente)."""
from __future__ import annotations

import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from ingest_medicoes import (
    PROFESSIONAL_COLUMNS,
    OperationalIdentityResolver,
    UnresolvedIdentityError,
    assert_operational_identities_resolved,
    read_measurements_sheet,
)
from test_negative_measurement_validation import normal_row

HEADERS = [
    "Número da Medição", "Projeto Referente", "Número do Documento", "Evidência", "Formato", "Quantidade",
    "Multiplicador", "Equivalente (A1 ou Horas)", "Medido (Horas)", "Item da QQP", "Valor Unitário", "Valor Bruto",
    "Valor Total", "Valor do Reajuste", "CICLO", "PROJETISTA", "REFERÊNCIA", "% EMISSÃO", "TIPO", "VALOR DE MEDIÇÃO",
]
FORNECEDOR_A = {"id": "p-a", "codigo": "FORNECEDOR ALFA TESTE", "nome": "FORNECEDOR ALFA TESTE"}
FORNECEDOR_B = {"id": "p-b", "codigo": "FORNECEDOR BETA TESTE", "nome": "FORNECEDOR BETA TESTE"}


def row_values(**overrides) -> list:
    base = normal_row(**overrides)
    return [base.get(header) for header in HEADERS]


def build_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(["PLANILHA SINTÉTICA DE TESTE"])  # título acima do cabeçalho (cabeçalho é localizado)
    ws.append(HEADERS)                            # linha 2
    projetista_col = HEADERS.index("PROJETISTA") + 1
    valor_col = HEADERS.index("Valor Bruto") + 1
    # Linhas 3–6: PROJETISTA mesclado (âncora na 3) — antes viravam "GRD-…" pela Evidência.
    for numero in range(4):
        ws.append(row_values(PROJETISTA=None, Evidência="GRD-T-SINT-2026-0001-0001", **{"Número do Documento": f"DOC-SINT-{numero}"}))
    ws.cell(3, projetista_col).value = "FORNECEDOR ALFA TESTE"
    ws.merge_cells(start_row=3, start_column=projetista_col, end_row=6, end_column=projetista_col)
    # Linha 7: valor numérico mesclado nunca é propagado (linha 8 fica sem Valor Bruto).
    ws.append(row_values(PROJETISTA="FORNECEDOR BETA TESTE"))
    ws.append(row_values(PROJETISTA="FORNECEDOR BETA TESTE", **{"Valor Bruto": None}))
    ws.merge_cells(start_row=7, start_column=valor_col, end_row=8, end_column=valor_col)
    # Linhas 9–11: medições sem PROJETISTA com texto de descrição / documento nas outras colunas.
    ws.append(row_values(PROJETISTA=None, Evidência="HORAS DE ESTUDO SINTETICAS PARA TESTE", **{"Número do Documento": "ORC-SINT-0001"}))
    ws.append(row_values(PROJETISTA=None, Evidência="TUBULAÇÃO"))
    ws.append(row_values(PROJETISTA="   ", Evidência="GRD-T-SINT-2026-0002"))
    wb.save(path)


def read_fixture() -> pd.DataFrame:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "fixture-identidades.xlsx"
        build_workbook(path)
        return read_measurements_sheet(path, "Documentos")


def preflight(df: pd.DataFrame, bm_aux: pd.DataFrame | None = None):
    r = OperationalIdentityResolver([FORNECEDOR_A, FORNECEDOR_B], [])
    return assert_operational_identities_resolved(
        df, pd.DataFrame(), bm_aux if bm_aux is not None else pd.DataFrame(), pd.DataFrame(), set(), {}, {},
        "2608", r, {}, "Documentos", "Base", "Documentos Auxiliares", None,
    )


def test_identity_comes_only_from_projetista() -> None:
    assert PROFESSIONAL_COLUMNS["nome"] == ["PROJETISTA"]


def test_merged_projetista_inherits_anchor_value_only_in_that_column() -> None:
    df = read_fixture()
    assert df.attrs["excel_row_numbers"][:6] == [3, 4, 5, 6, 7, 8]
    assert df["PROJETISTA"].tolist()[:4] == ["FORNECEDOR ALFA TESTE"] * 4
    assert df["PROJETISTA"].tolist()[4:6] == ["FORNECEDOR BETA TESTE"] * 2
    valores = df["Valor Bruto"].tolist()
    assert valores[4] == 571.41 and pd.isna(valores[5])  # valor mesclado não é duplicado


def test_rows_without_projetista_block_and_never_become_identities() -> None:
    df = read_fixture()
    try:
        preflight(df)
    except UnresolvedIdentityError as error:
        assert [d["status"] for d in error.details] == ["SEM_PROJETISTA"]
        detalhe = error.details[0]
        assert detalhe["origem"] == "Documentos" and detalhe["coluna"] == "PROJETISTA"
        assert detalhe["ocorrencias"] == 3 and detalhe["linhas"] == [9, 10, 11]
        assert detalhe["valor"] == "" and detalhe["candidatos"] == []
        assert detalhe["exemplos"][0] == {"linha": 9, "numeroDocumento": "ORC-SINT-0001", "evidencia": "HORAS DE ESTUDO SINTETICAS PARA TESTE"}
        valores = {d["valor"] for d in error.details}
        assert not any(v.startswith(("GRD", "ORC", "HORAS", "TUBULA")) for v in valores if v)
        mensagem = str(error)
        assert "3 linha(s) possuem dados de medição, mas não têm PROJETISTA válido" in mensagem and "Nenhum dado foi alterado" in mensagem
        assert "descarte conscientemente" in mensagem
    else:
        raise AssertionError("Linhas de medição sem PROJETISTA deveriam bloquear a importação.")


def test_missing_projetista_blocks_while_unknown_name_stays_pending() -> None:
    """Linha com dados sem PROJETISTA bloqueia; o nome desconhecido ao lado dela NÃO é erro (pendente)."""
    rows = [normal_row(PROJETISTA="NOME SINTETICO NOVO"), normal_row(PROJETISTA=None), normal_row(PROJETISTA="NOME SINTETICO NOVO")]
    df = pd.DataFrame(rows)
    df.attrs["excel_row_numbers"] = [10, 11, 12]
    try:
        preflight(df)
    except UnresolvedIdentityError as error:
        assert [d["status"] for d in error.details] == ["SEM_PROJETISTA"]
        assert error.details[0]["linhas"] == [11]
        assert len(error.details[0]["chaves"]) == 1  # chave estável para o descarte consciente
        assert '"NOME SINTETICO NOVO"' not in str(error)
    else:
        raise AssertionError("Linha sem PROJETISTA deveria bloquear.")
    summary = preflight(df.iloc[[0, 2]])
    assert [p["valor"] for p in summary["identidades_pendentes"]] == ["NOME SINTETICO NOVO"]
    assert summary["plano"]["pendentes"]["NOME SINTETICO NOVO"]["coluna"] == "PROJETISTA"


def test_fully_resolved_sheet_passes() -> None:
    df = read_fixture().iloc[:6]
    summary = preflight(df)
    assert summary["identidades_resolvidas"] == 2 and summary["identidades_por_via"] == {"CODIGO": 2}


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok {name}")
